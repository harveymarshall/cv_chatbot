import os
from langchain.memory import ConversationBufferMemory
from app.services.pdf_to_text import extract_pdf_text
from app.services.docx_to_pdf import docx_to_pdf
from app.services.add_msg_to_history import add_msg_to_history


def test_extract_pdf_text(tmp_path):
    # Create a simple PDF file for testing
    from fpdf import FPDF
    pdf_path = tmp_path / "test.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Hello PDF!", ln=True)
    pdf.output(str(pdf_path))

    # Test extraction
    text = extract_pdf_text(str(pdf_path))
    assert "Hello PDF!" in text

def test_docx_to_pdf(tmp_path):
    # Create a simple DOCX file for testing
    from docx import Document
    docx_path = tmp_path / "test.docx"
    pdf_path = tmp_path / "test.pdf"
    doc = Document()
    doc.add_paragraph("Hello DOCX!")
    doc.save(str(docx_path))

    # Convert to PDF
    docx_to_pdf(str(docx_path), str(pdf_path))
    assert os.path.exists(pdf_path)
    # Optionally, check PDF content
    from app.services.pdf_to_text import extract_pdf_text
    text = extract_pdf_text(str(pdf_path))
    assert "Hello DOCX!" in text

def test_add_msg_to_history():
    history = [{"role": "user", "content": "hello_world"}]
    memory = ConversationBufferMemory(return_messages=True)
    add_msg_to_history(memory, history)
    assert memory.chat_memory.messages[0].content == "hello_world"

